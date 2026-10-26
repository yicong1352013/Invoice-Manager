"""Pure spreadsheet rendering and the default template seed."""

from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.models.workflow import ExportTemplate
from app.schemas.export_template import ExcelTemplateConfig, FIELD_CATALOG


DEFAULT_CONFIG = ExcelTemplateConfig(columns=[
    {"field": field["key"], "title": field["label"], "format": field["type"],
     "sum": field["summable"], "width": 24 if field["type"] == "text" else 18,
     "align": "right" if field["summable"] else "left"}
    for field in FIELD_CATALOG[:17]
]).model_dump()


def default_config():
    return deepcopy(DEFAULT_CONFIG)


def sample_invoices():
    return [{
        "invoice_number": "00123456789012345678", "issue_date": "2026-10-24",
        "invoice_type": "电子发票（普通发票）", "buyer_name": "示例购买方",
        "buyer_tax_id": "001234567890123456", "seller_name": "示例供应商",
        "seller_tax_id": "009876543210987654", "item_name": "办公用品",
        "amount": "1000.00", "tax_amount": "130.00", "total_with_tax": "1130.00",
        "tax_rate": "13%", "remark": "内置样例数据", "status": "已确认",
        "owner": "样例用户", "file_name": "示例发票.pdf", "created_at": "2026-10-24T09:00:00",
    }]


def _text(value):
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))[:32767]


def _value(invoice, column):
    value = invoice.get(column.field) if isinstance(invoice, dict) else getattr(invoice, column.field, None)
    if value is None or value == "":
        return None
    if hasattr(value, "value"):
        value = value.value
    if column.format in ("money", "number"):
        try:
            value = Decimal(str(value))
        except InvalidOperation as exc:
            raise ValueError(f"{column.title} 不是有效数值") from exc
        if not value.is_finite():
            raise ValueError(f"{column.title} 不是有限数值")
        return value
    if column.format in ("date", "datetime"):
        if not isinstance(value, (date, datetime)):
            value = datetime.fromisoformat(str(value))
        if column.format == "date":
            return value.date() if isinstance(value, datetime) else value
        if not isinstance(value, datetime):
            value = datetime.combine(value, datetime.min.time())
        return value.replace(tzinfo=None)
    return _text(value)


def _normalized(invoices, config):
    config = config if isinstance(config, ExcelTemplateConfig) else ExcelTemplateConfig.model_validate(config)
    rows = [[_value(invoice, column) for column in config.columns] for invoice in invoices]
    totals = [sum((row[index] for row in rows if row[index] is not None), Decimal(0))
              if column.sum else None for index, column in enumerate(config.columns)]
    return config, rows, totals


def _display(value, column):
    if value is None:
        return None
    if column.format == "money":
        return f"{value:,.2f}"
    if column.format == "number":
        return f"{value:,.4f}".rstrip("0").rstrip(".")
    if column.format == "date":
        return value.strftime("%Y-%m-%d")
    if column.format == "datetime":
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def preview_excel(invoices, config):
    config, rows, totals = _normalized(invoices, config)
    return {"columns": [column.model_dump() for column in config.columns],
            "rows": [[_display(value, col) for value, col in zip(row, config.columns)] for row in rows],
            "totals": [_display(value, col) for value, col in zip(totals, config.columns)],
            "count": len(rows), "title": config.title, "sheet_name": config.sheet_name}


def build_excel_export(invoices, config=None) -> BytesIO:
    config, rows, totals = _normalized(invoices, config if config is not None else DEFAULT_CONFIG)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = config.sheet_name
    header_row = 2 if config.title else 1
    border = Border(**{edge: Side(style="thin", color="D9D9D9")
                       for edge in ("left", "right", "top", "bottom")})

    def write(row, col, value, bold=False):
        cell = sheet.cell(row, col, value)
        # Force all user-controlled strings to text: '=' must never become a formula.
        if isinstance(value, str):
            cell.data_type = "s"
        cell.font = Font(size=config.style.font_size, bold=bold)
        if config.style.borders:
            cell.border = border
        return cell

    if config.title:
        write(1, 1, _text(config.title), True)
        if len(config.columns) > 1:
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(config.columns))
    formats = {"text": "@", "date": "yyyy-mm-dd", "datetime": "yyyy-mm-dd hh:mm:ss",
               "money": "#,##0.00", "number": "#,##0.####"}
    for index, column in enumerate(config.columns, 1):
        cell = write(header_row, index, _text(column.title), True)
        cell.fill = PatternFill("solid", fgColor=config.style.header_color[1:])
        cell.font = Font(size=config.style.font_size, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    for row_index, row in enumerate(rows, header_row + 1):
        for col_index, (value, column) in enumerate(zip(row, config.columns), 1):
            cell = write(row_index, col_index, value)
            cell.number_format = formats[column.format]
            cell.alignment = Alignment(horizontal=column.align)
            if config.style.striped and (row_index - header_row) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FC")
    if any(column.sum for column in config.columns):
        row_index = header_row + len(rows) + 1
        for index, (value, column) in enumerate(zip(totals, config.columns), 1):
            cell = write(row_index, index, value, True)
            cell.number_format = formats[column.format]
        label_index = next((i for i, c in enumerate(config.columns, 1) if not c.sum), None)
        if label_index:
            write(row_index, label_index, "合计", True)
    if config.show_count:
        count_row = header_row + len(rows) + 1 + int(any(c.sum for c in config.columns))
        write(count_row, 1, f"发票数量：{len(rows)} 张", True)
    if config.freeze_header:
        sheet.freeze_panes = f"A{header_row + 1}"
    if config.auto_filter:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(config.columns))}{header_row + len(rows)}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_export_filename(config, run_id="", run_name="", extension="xlsx", now=None):
    config = config if isinstance(config, ExcelTemplateConfig) else ExcelTemplateConfig.model_validate(config)
    name = config.filename_pattern
    values = {"date": (now or datetime.now()).strftime("%Y%m%d"), "run_id": str(run_id), "run_name": str(run_name)}
    name = re.sub(r"\{(date|run_id|run_name)\}", lambda match: values[match.group(1)], name)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")[:100] or "发票导出"
    if name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)], *[f"LPT{i}" for i in range(1, 10)]}:
        name = "_" + name
    if extension not in {"xlsx", "docx", "pdf", "zip"}:
        raise ValueError("不支持的导出扩展名")
    return f"{name}.{extension}"


async def ensure_default_template(db):
    existing = await db.scalar(select(ExportTemplate).order_by(ExportTemplate.id).limit(1))
    if existing:
        return existing
    template = ExportTemplate(name="默认发票明细表", description="兼容原有 17 列发票导出",
                              draft_config=default_config(), published_config=default_config(), version=1, enabled=True)
    db.add(template)
    await db.flush()
    return template
