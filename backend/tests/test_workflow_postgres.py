"""Opt-in PostgreSQL checks in a disposable schema, never in business tables."""

import asyncio
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import subprocess
import sys
import uuid
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy import select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.config import get_settings
from app.database import Base
from app.models.invoice import Invoice, InvoiceStatus, OcrResult
from app.models.user import User, UserRole
from app.models.workflow import Delivery, ExportArtifact, ExportTemplate, NotificationChannel, WorkflowDefinition, WorkflowRun, WorkflowRunItem
from app.services.excel_template_service import default_config


pytestmark = pytest.mark.skipif(os.environ.get("WORKFLOW_POSTGRES_TESTS") != "1", reason="需要显式启用本地隔离 PostgreSQL 检查")


@pytest.mark.asyncio
async def test_process_restart_recovers_unknown_delivery_and_preserves_business_isolation(monkeypatch):
    url = make_url(get_settings().database_url)
    assert url.host in {"localhost", "127.0.0.1", "::1"}, "仅允许本地 PostgreSQL"
    schema = "workflow_test_" + uuid.uuid4().hex
    worker_lock = uuid.uuid4().int % (2 ** 63)
    control = create_async_engine(url, echo=False)
    isolated = create_async_engine(url, echo=False, connect_args={"server_settings": {"search_path": schema}})
    maker = async_sessionmaker(isolated, expire_on_commit=False)
    try:
        async with control.begin() as connection:
            await connection.execute(text(f'CREATE SCHEMA "{schema}"'))
        async with isolated.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
            assert (await connection.execute(text("SELECT current_schema()"))).scalar() == schema
        async with maker() as db:
            db.add_all([
                User(username="workflow-test", role=UserRole.USER, password_hash="test-only", enabled=True),
                WorkflowDefinition(id=1, name="isolated", config={}, enabled=True),
                NotificationChannel(id=1, name="NEVER SEND", kind="email", config={}, enabled=True),
                Invoice(id=1, file_name="isolated.pdf", file_type="pdf", file_data=b"isolated", owner="workflow-test", status=InvoiceStatus.CONFIRMED),
            ])
            await db.flush()
            db.add(WorkflowRun(id=1, workflow_id=1, name="isolated", owner="workflow-test", state="sending",
                               frozen_at=datetime.utcnow(), config_snapshot={"formats": ["excel"],
                               "deliveries": [{"channel_id": 1, "targets": ["nobody@example.invalid"]}]}, upload_manifest={}))
            await db.flush()
            db.add_all([
                WorkflowRunItem(run_id=1, invoice_id=1, state="approved", data_snapshot={"total_with_tax": "0"}),
                ExportArtifact(run_id=1, kind="excel", file_name="same.xlsx", file_data=b"immutable bytes",
                               size=15, sha256="0" * 64, content_type="application/octet-stream"),
                Delivery(run_id=1, channel_id=1, target="nobody@example.invalid", part="success:envelope",
                         state="inflight", attempts=[{"at": "before-crash"}]),
            ])
            await db.commit()

        # Two independent database connections must not both own the worker.
        async with isolated.connect() as first, isolated.connect() as second:
            assert (await first.execute(text("SELECT pg_try_advisory_lock(:key)"), {"key": worker_lock})).scalar()
            assert not (await second.execute(text("SELECT pg_try_advisory_lock(:key)"), {"key": worker_lock})).scalar()
            await first.execute(text("SELECT pg_advisory_unlock(:key)"), {"key": worker_lock})

        # A fresh Python process uses only this schema; any accidental outbound
        # dispatch fails locally. Recovering an inflight ledger must not dispatch.
        program = """
import asyncio, os
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from app.config import get_settings
from app import workflow_worker
from app.services import notification_service
async def forbidden(*args, **kwargs):
    raise AssertionError('recovery must never send')
async def main():
    engine = create_async_engine(get_settings().database_url, echo=False,
        connect_args={'server_settings': {'search_path': os.environ['WORKFLOW_TEST_SCHEMA']}})
    workflow_worker.engine = engine
    workflow_worker.WORKER_LOCK = int(os.environ['WORKFLOW_TEST_LOCK'])
    workflow_worker.async_session_maker = async_sessionmaker(engine, expire_on_commit=False)
    notification_service.deliver = forbidden
    try:
        await workflow_worker.run_worker(once=True)
    finally:
        await engine.dispose()
asyncio.run(main())
"""
        environment = {**os.environ, "WORKFLOW_TEST_SCHEMA": schema, "WORKFLOW_TEST_LOCK": str(worker_lock), "DEBUG": "false"}
        for _ in range(2):
            result = await asyncio.to_thread(subprocess.run, [sys.executable, "-c", program],
                cwd=str(Path(__file__).resolve().parents[1]), env=environment, capture_output=True, text=True, timeout=30)
            assert result.returncode == 0, result.stderr[-2000:]
        async with maker() as db:
            run = await db.get(WorkflowRun, 1)
            delivery = (await db.execute(select(Delivery))).scalar_one()
            assert run.state == "delivery_failed"
            assert run.lease_owner is None and run.lease_until is None
            assert delivery.state == "unknown" and len(delivery.attempts) == 1
            assert (await db.execute(select(ExportArtifact))).scalar_one().file_data == b"immutable bytes"

        from app import database
        from app.main import _recover_stuck_processing
        monkeypatch.setattr(database, "async_session_maker", maker)
        async with maker() as db:
            (await db.get(Invoice, 1)).status = InvoiceStatus.PROCESSING
            db.add(Invoice(id=2, file_name="legacy.pdf", file_type="pdf", file_data=b"legacy", owner="workflow-test", status=InvoiceStatus.PROCESSING))
            await db.commit()
        await _recover_stuck_processing()
        async with maker() as db:
            assert (await db.get(Invoice, 1)).status == InvoiceStatus.PROCESSING
            assert (await db.get(Invoice, 2)).status == InvoiceStatus.FAILED

        # Exercise the complete executor on actual PostgreSQL types/transactions.
        # External OCR, AI and notification operations are isolated fakes.
        from app.services import workflow_service, notification_service
        from openpyxl import load_workbook
        calls = []

        async def recognize(invoice_id, db):
            (await db.get(Invoice, invoice_id)).status = InvoiceStatus.CONFIRMED
            await db.commit()
            return True

        async def verify(db, invoice):
            return {"verdict": {"overall": "pass", "fields": []}}

        async def accepted(*args, **kwargs):
            calls.append(args[1])
            return {"accepted": True, "provider": "isolated-fake"}

        monkeypatch.setattr(workflow_service, "process_invoice", recognize)
        monkeypatch.setattr(workflow_service, "verify_invoice", verify)
        monkeypatch.setattr(notification_service, "deliver", accepted)
        async with maker() as db:
            db.add(Invoice(id=3, file_name="e2e.pdf", file_type="pdf", file_data=b"e2e", owner="workflow-test",
                           invoice_number="00123456789012345678", issue_date=date(2026, 10, 24),
                           buyer_name="测试买方", buyer_tax_id="001234567890123", seller_name="测试卖方",
                           seller_tax_id="009876543210987", item_name="测试费用", amount=Decimal("0"),
                           tax_amount=Decimal("0"), total_with_tax=Decimal("0"), tax_rate="免税"))
            db.add(WorkflowRun(id=2, workflow_id=1, name="PostgreSQL 闭环", owner="workflow-test", state="queued",
                               config_snapshot={"formats": ["excel"], "template": default_config(),
                               "deliveries": [{"channel_id": 1, "targets": ["nobody@example.invalid"]}]},
                               upload_manifest={"invoices": [{"id": "e2e", "name": "e2e.pdf"}]}))
            await db.flush()
            db.add(WorkflowRunItem(run_id=2, invoice_id=3, client_file_id="e2e", state="pending"))
            await db.commit()
        async with maker() as db:
            await workflow_service.process_run(db, await db.get(WorkflowRun, 2))
        async with maker() as db:
            assert (await db.get(WorkflowRun, 2)).state == "completed"
            artifact = (await db.execute(select(ExportArtifact).where(ExportArtifact.run_id == 2))).scalar_one()
            sheet = load_workbook(BytesIO(artifact.file_data)).active
            assert sheet.cell(2, 1).value == "00123456789012345678"
            assert sheet.cell(2, 11).value == 0
            assert calls == ["nobody@example.invalid"]

        from app.routers import workflows
        from app.schemas.workflow import WorkflowRunCreate

        # Cancellation must survive both a stale worker and a racing transition.
        async with maker() as db:
            for run_id in (3, 4, 5):
                db.add(WorkflowRun(id=run_id, workflow_id=1, name="cancel isolation", owner="workflow-test", state="queued",
                                   config_snapshot={"formats": ["excel"], "template": default_config(),
                                   "deliveries": [{"channel_id": 1, "targets": ["nobody@example.invalid"]}]}))
                await db.flush()
                db.add(WorkflowRunItem(run_id=run_id, invoice_id=3, state="recognized"))
            await db.commit()
        async with maker() as worker, maker() as cancelling:
            stale = await worker.get(WorkflowRun, 3)
            actor = (await cancelling.execute(select(User).where(User.username == "workflow-test"))).scalar_one()
            assert (await workflows.cancel_run(3, actor, cancelling))["state"] == "cancelled"
            assert stale.state == "queued"
            await workflow_service.process_run(worker, stale)
            assert stale.state == "cancelled"

        transition_locked, resume_transition = asyncio.Event(), asyncio.Event()
        original_cancelled = workflow_service._cancelled

        async def pause_transition(db, run, lock=False):
            cancelled = await original_cancelled(db, run, lock=lock)
            if run.id == 4 and lock and not transition_locked.is_set():
                transition_locked.set()
                await resume_transition.wait()
            return cancelled

        monkeypatch.setattr(workflow_service, "_cancelled", pause_transition)
        async with maker() as worker, maker() as cancelling:
            actor = (await cancelling.execute(select(User).where(User.username == "workflow-test"))).scalar_one()
            task = asyncio.create_task(workflow_service.process_run(worker, await worker.get(WorkflowRun, 4)))
            await asyncio.wait_for(transition_locked.wait(), timeout=5)
            cancel = asyncio.create_task(workflows.cancel_run(4, actor, cancelling))
            try:
                await asyncio.sleep(0.05)
                assert not cancel.done(), "Cancel must serialize with the locked state transition"
            finally:
                resume_transition.set()
                results = await asyncio.wait_for(asyncio.gather(task, cancel), timeout=5)
            assert results[1]["state"] == "cancelled"
        monkeypatch.setattr(workflow_service, "_cancelled", original_cancelled)

        # No run row lock may span an external AI call and block cancellation.
        async def cancel_during_verify(db, invoice):
            async with maker() as cancelling:
                actor = (await cancelling.execute(select(User).where(User.username == "workflow-test"))).scalar_one()
                result = await asyncio.wait_for(workflows.cancel_run(5, actor, cancelling), timeout=5)
                assert result["state"] == "cancelled"
            return {"verdict": {"overall": "pass", "fields": []}}

        monkeypatch.setattr(workflow_service, "verify_invoice", cancel_during_verify)
        async with maker() as worker:
            await workflow_service.process_run(worker, await worker.get(WorkflowRun, 5))
        async with maker() as db:
            for run_id in (3, 4, 5):
                assert (await db.get(WorkflowRun, run_id)).state == "cancelled"
                assert not (await db.execute(select(ExportArtifact).where(ExportArtifact.run_id == run_id))).scalars().all()
                assert not (await db.execute(select(Delivery).where(Delivery.run_id == run_id))).scalars().all()
            assert calls == ["nobody@example.invalid"]

        # Two simultaneous list actions must serialize on the selected invoice.
        async def provider(*args):
            return SimpleNamespace(enabled=True, api_key="isolated-only", model="test")
        monkeypatch.setattr(workflows, "get_active_provider_row", provider)
        async with maker() as db:
            template = ExportTemplate(name="selection", draft_config=default_config(), published_config=default_config(), version=1, enabled=True)
            db.add(template)
            await db.flush()
            db.add(WorkflowDefinition(id=2, name="selection", enabled=True, config={
                "formats": ["excel"], "template_id": template.id,
                "deliveries": [{"channel_id": 1, "targets": ["nobody@example.invalid"]}],
            }))
            db.add(OcrResult(invoice_id=3, raw_text="existing recognition"))
            await db.execute(text("SELECT setval(pg_get_serial_sequence('workflow_runs', 'id'), (SELECT max(id) FROM workflow_runs))"))
            await db.commit()
        async def select_existing():
            async with maker() as db:
                actor = (await db.execute(select(User).where(User.username == "workflow-test"))).scalar_one()
                return await workflows.create_run(WorkflowRunCreate(workflow_id=2, name="concurrent selection", invoice_ids=[3]), actor, db)
        results = await asyncio.gather(select_existing(), select_existing(), return_exceptions=True)
        accepted = [result for result in results if isinstance(result, dict)]
        rejected = [result for result in results if isinstance(result, HTTPException)]
        assert len(accepted) == 1 and accepted[0]["state"] == "queued", results
        assert len(rejected) == 1 and rejected[0].status_code == 409, results
        async with maker() as db:
            item = (await db.execute(select(WorkflowRunItem).where(WorkflowRunItem.run_id == accepted[0]["id"]))).scalar_one()
            assert item.invoice_id == 3 and item.state == "recognized"
    finally:
        await isolated.dispose()
        async with control.begin() as connection:
            # The generated schema name is fixed-prefix + UUID hex, never a path
            # or user input. No business schema/table is dropped.
            assert schema.startswith("workflow_test_") and len(schema) == 46
            await connection.execute(text(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE'))
        await control.dispose()
