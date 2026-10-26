"""A workflow never recognizes an incomplete upload or duplicates a retried file."""

import json
from io import BytesIO
from urllib.parse import quote
from unittest.mock import AsyncMock

from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.auth import get_current_user
from app.database import Base, get_db
from app.models.attachment import Attachment
from app.models.invoice import Invoice, InvoiceStatus
from app.models.user import User, UserRole
from app.models.workflow import ExportTemplate, WorkflowDefinition, WorkflowRun, WorkflowRunItem
from app.routers import invoices


class AsyncAdapter:
    def __init__(self, session):
        self.session = session

    async def execute(self, query):
        return self.session.execute(query)

    async def get(self, model, key):
        return self.session.get(model, key)

    def add(self, value):
        self.session.add(value)

    async def flush(self):
        self.session.flush()

    async def commit(self):
        self.session.commit()

    async def rollback(self):
        self.session.rollback()

    async def refresh(self, value):
        self.session.refresh(value)

    async def delete(self, value):
        self.session.delete(value)


@pytest.fixture
def api(monkeypatch):
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        user = User(username="alice", role=UserRole.USER, enabled=True, password_hash="unused")
        definition = WorkflowDefinition(id=1, name="test", config={}, enabled=True)
        session.add_all([user, definition])
        session.flush()
        run = WorkflowRun(id=1, workflow_id=1, name="test", owner="alice", state="draft",
                          config_snapshot={}, upload_manifest={"invoices": [{"id": "inv-a", "name": "a.pdf"}],
                          "attachments": [{"id": "att-a", "name": "contract.pdf"}]})
        session.add(run)
        session.commit()
        app = FastAPI()
        app.include_router(invoices.router, prefix="/api/invoices")
        app.dependency_overrides[get_current_user] = lambda: user
        app.dependency_overrides[get_db] = lambda: AsyncAdapter(session)
        monkeypatch.setattr(invoices.limiter, "enabled", False)
        background = AsyncMock()
        monkeypatch.setattr(invoices, "process_invoice_background", background)
        with TestClient(app) as client:
            yield client, session, run, user, background
    engine.dispose()


def upload(client, key="inv-a", content=b"invoice", filename="a.pdf", run_id=1):
    return client.post("/api/invoices/upload", data={"workflow_run_id": run_id,
                       "client_file_ids": json.dumps([key])}, files=[("files", (filename, content, "application/pdf"))])


def test_workflow_upload_is_idempotent_and_never_starts_ocr(api):
    client, session, _, _, background = api
    first = upload(client)
    assert first.status_code == 200 and first.json()[0]["status"] == "success"
    again = upload(client)
    assert again.json()[0]["id"] == first.json()[0]["id"]
    assert len(session.scalars(select(Invoice)).all()) == 1
    assert len(session.scalars(select(WorkflowRunItem)).all()) == 1
    assert session.scalars(select(Invoice)).one().status == InvoiceStatus.UPLOADED
    background.assert_not_called()


def test_retry_cannot_replace_a_manifest_file_with_different_content(api):
    client, session, run, _, _ = api
    upload(client)
    response = upload(client, content=b"different invoice")
    assert response.json()[0]["status"] == "error"
    assert "inv-a" in run.upload_errors
    assert session.scalars(select(Invoice)).one().file_data == b"invoice"
    assert upload(client).json()[0]["status"] == "success"
    assert run.upload_errors == {}


def test_foreign_sealed_or_mismatched_manifest_upload_is_rejected(api):
    client, _, run, user, _ = api
    user.username = "bob"
    assert upload(client).status_code == 404
    user.username = "alice"
    assert upload(client, key="unknown").status_code == 409
    assert upload(client, filename="different.pdf").status_code == 409
    run.state = "queued"
    assert upload(client).status_code == 409


def test_attachment_must_belong_to_batch_and_retry_is_idempotent(api):
    client, session, run, _, background = api
    invoice_id = upload(client).json()[0]["id"]
    unrelated = Invoice(file_name="other.pdf", file_type="pdf", file_data=b"other", owner="alice")
    session.add(unrelated)
    session.commit()
    data = {"workflow_run_id": 1, "client_file_ids": '["att-a"]'}
    files = [("files", ("contract.pdf", b"contract", "application/pdf"))]
    assert client.post(f"/api/invoices/{unrelated.id}/attachments", data=data, files=files).status_code == 409
    url = f"/api/invoices/{invoice_id}/attachments"
    first = client.post(url, data=data, files=files)
    assert first.status_code == 200 and first.json()[0]["status"] == "success"
    assert client.post(url, data=data, files=files).json()[0]["id"] == first.json()[0]["id"]
    assert len(session.scalars(select(Attachment)).all()) == 1
    assert run.uploaded_attachments == {"att-a": first.json()[0]["id"]}
    background.assert_not_called()


def test_original_upload_still_schedules_recognition(api):
    client, session, _, _, background = api
    result = client.post("/api/invoices/upload", files=[("files", ("normal.pdf", b"invoice", "application/pdf"))])
    assert result.status_code == 200 and result.json()[0]["status"] == "success"
    background.assert_awaited_once_with(result.json()[0]["id"])
    assert session.scalars(select(WorkflowRunItem)).all() == []


def test_linked_invoices_cannot_be_deleted_and_active_run_blocks_manual_reprocess(api):
    client, _, run, _, _ = api
    invoice_id = upload(client).json()[0]["id"]
    assert client.delete(f"/api/invoices/{invoice_id}").status_code == 409
    response = client.post("/api/invoices/batch-delete", json={"invoice_ids": [invoice_id]})
    assert response.status_code == 409
    assert client.post(f"/api/invoices/{invoice_id}/process").status_code == 409
    run.state = "recognizing"
    assert client.post("/api/invoices/batch-reprocess", json={"invoice_ids": [invoice_id]}).status_code == 409


def test_existing_excel_endpoint_preserves_default_and_accepts_published_template(api):
    from openpyxl import load_workbook
    from app.services.excel_template_service import default_config
    client, session, _, _, _ = api
    identifier = upload(client).json()[0]["id"]
    invoice = session.get(Invoice, identifier)
    invoice.invoice_number, invoice.total_with_tax = "00123456789012345678", 0
    config = default_config()
    config.update(title="自定义报表", filename_pattern="报销_{date}")
    template = ExportTemplate(name="custom", draft_config=config, published_config=config, version=1, enabled=True)
    session.add(template)
    session.commit()
    url = f"/api/invoices/export/excel?invoice_ids={identifier}"
    standard = client.get(url)
    assert standard.status_code == 200 and quote("发票导出.xlsx") in standard.headers["content-disposition"]
    sheet = load_workbook(BytesIO(standard.content)).active
    assert sheet.cell(2, 1).value == "00123456789012345678" and sheet.cell(2, 11).value == 0
    custom = client.get(url + f"&template_id={template.id}")
    assert custom.status_code == 200 and quote("报销_") in custom.headers["content-disposition"]
    assert load_workbook(BytesIO(custom.content)).active.cell(1, 1).value == "自定义报表"
    template.enabled = False
    session.commit()
    assert client.get(url + f"&template_id={template.id}").status_code == 404


def test_failed_manual_recognition_does_not_leave_invoice_processing_forever(api, monkeypatch):
    from app.services import invoice_service
    client, session, _, _, _ = api
    result = client.post("/api/invoices/upload", files=[("files", ("normal.pdf", b"invoice", "application/pdf"))])
    identifier = result.json()[0]["id"]
    monkeypatch.setattr(invoice_service, "process_invoice", AsyncMock(return_value=False))
    assert client.post(f"/api/invoices/{identifier}/process").status_code == 500
    assert session.get(Invoice, identifier).status == InvoiceStatus.FAILED
