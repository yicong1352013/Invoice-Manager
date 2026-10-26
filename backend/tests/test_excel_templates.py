"""Business invariants for template exports, publication and data ownership."""

from datetime import date, datetime
from decimal import Decimal

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from pydantic import ValidationError
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.auth import get_current_user
from app.database import Base, get_db
from app.models.invoice import Invoice, InvoiceStatus
from app.models.user import User, UserRole
from app.models.workflow import ExportTemplate
from app.routers.export_templates import router
from app.schemas.export_template import ExcelTemplateConfig
from app.services.excel_template_service import (
    build_excel_export, build_export_filename, default_config, preview_excel,
)


def test_excel_preserves_identifier_zero_blank_and_literal_formula_text():
    config = default_config()
    rows = [{"invoice_number": "001234567890123456789", "issue_date": date(2026, 10, 24),
             "buyer_tax_id": "000987654321012345678", "seller_name": "=HYPERLINK(\"evil\")",
             "amount": Decimal("0"), "tax_amount": None, "total_with_tax": Decimal("0.00")}]
    worksheet = load_workbook(build_excel_export(rows, config)).active
    assert worksheet["A2"].value == "001234567890123456789"
    assert worksheet["E2"].value == "000987654321012345678"
    assert worksheet["F2"].data_type == "s"
    assert worksheet["F2"].value.startswith("=HYPERLINK")
    assert worksheet["I2"].value == 0
    assert worksheet["I2"].number_format == "#,##0.00"
    assert worksheet["J2"].value is None
    assert worksheet["B2"].value == datetime(2026, 10, 24)
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:Q2"
    preview = preview_excel(rows, config)
    assert preview["rows"][0][8:11] == ["0.00", None, "0.00"]
    assert preview["rows"][0][1] == "2026-10-24"


def test_totals_follow_column_reordering_without_overwriting_first_money_column():
    config = default_config()
    config["columns"] = [config["columns"][10], config["columns"][8]]
    rows = [{"total_with_tax": Decimal("0.10"), "amount": Decimal("0.10")},
            {"total_with_tax": Decimal("0.20"), "amount": None}]
    worksheet = load_workbook(build_excel_export(rows, config)).active
    assert worksheet["A4"].value == 0.3
    assert worksheet["B4"].value == 0.1
    assert worksheet["A5"].value == "发票数量：2 张"
    assert preview_excel(rows, config)["totals"] == ["0.30", "0.10"]


@pytest.mark.parametrize("mutation", [
    lambda c: c.update(columns=[]),
    lambda c: c["columns"].append(c["columns"][0]),
    lambda c: c["columns"][0].update(format="number"),
    lambda c: c["columns"][0].update(sum=True),
    lambda c: c["columns"][0].update(field="file_data"),
    lambda c: c["columns"][1].update(title=c["columns"][0]["title"]),
    lambda c: c.update(sheet_name="bad/name"),
    lambda c: c.update(filename_pattern="../secret"),
    lambda c: c.update(filename_pattern="{unknown}"),
])
def test_invalid_templates_cannot_publish_or_export(mutation):
    config = default_config()
    mutation(config)
    with pytest.raises(ValidationError):
        build_excel_export([], config)


def test_titles_remain_literal_and_sample_uses_same_style():
    config = default_config()
    config["title"] = "=1+1"
    config["columns"][0]["title"] = "=2+2"
    worksheet = load_workbook(build_excel_export([], config)).active
    assert worksheet["A1"].data_type == "s"
    assert worksheet["A2"].data_type == "s"
    assert worksheet.freeze_panes == "A3"
    assert worksheet["A2"].fill.fgColor.rgb == "001677FF"


def test_filenames_are_bounded_and_cannot_escape_download_directory():
    config = default_config()
    config["filename_pattern"] = "{run_name}_{run_id}_{date}"
    filename = build_export_filename(config, 42, "../../name", now=datetime(2026, 10, 24))
    assert filename.endswith("_42_20261024.xlsx")
    assert "/" not in filename and "\\" not in filename
    assert len(build_export_filename(config, 42, "长" * 300)) <= 106
    config["filename_pattern"] = "CON"
    assert build_export_filename(config) == "_CON.xlsx"


class AsyncSessionAdapter:
    """Run endpoint SQL against a real isolated SQLite DB without a new driver dependency."""

    def __init__(self, session):
        self.session = session

    async def execute(self, statement):
        return self.session.execute(statement)

    async def scalar(self, statement):
        return self.session.scalar(statement)

    def add(self, value):
        self.session.add(value)

    async def commit(self):
        self.session.commit()

    async def refresh(self, value):
        self.session.refresh(value)

    async def delete(self, value):
        self.session.delete(value)


@pytest.fixture
def api():
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add_all([
            Invoice(id=1, file_name="mine.pdf", file_type="pdf", file_data=b"x", owner="alice",
                    invoice_number="mine", status=InvoiceStatus.CONFIRMED),
            Invoice(id=2, file_name="theirs.pdf", file_type="pdf", file_data=b"x", owner="bob",
                    invoice_number="private", status=InvoiceStatus.CONFIRMED),
        ])
        session.commit()
        app = FastAPI()
        app.include_router(router, prefix="/api/export-templates")
        current = {"user": User(username="alice", role=UserRole.USER, enabled=True)}
        app.dependency_overrides[get_current_user] = lambda: current["user"]
        app.dependency_overrides[get_db] = lambda: AsyncSessionAdapter(session)
        with TestClient(app) as client:
            yield client, current, session
    engine.dispose()


def test_preview_never_exposes_other_owner_invoice(api):
    client, _, _ = api
    body = {"config": default_config(), "invoice_ids": [1]}
    response = client.post("/api/export-templates/preview", json=body)
    assert response.status_code == 200
    assert response.json()["rows"][0][0] == "mine"
    body["invoice_ids"] = [1, 2]
    assert client.post("/api/export-templates/preview", json=body).status_code == 404
    assert client.post("/api/export-templates/sample", json=body).status_code == 404
    assert client.post("/api/export-templates", json={"name": "no", "config": default_config()}).status_code == 403


def test_publish_versions_keep_existing_snapshot_and_hide_drafts(api):
    client, current, session = api
    current["user"].role = UserRole.ADMIN
    response = client.post("/api/export-templates", json={"name": "财务汇总", "config": default_config()})
    assert response.status_code == 201
    identifier = response.json()["id"]
    url = f"/api/export-templates/{identifier}"
    current["user"].role = UserRole.USER
    assert client.get("/api/export-templates").json() == []
    current["user"].role = UserRole.ADMIN
    published = client.post(url + "/publish", json={"expected_version": 0})
    assert published.status_code == 200
    snapshot = published.json()["published_config"]
    changed = default_config()
    changed["title"] = "下一版"
    assert client.put(url, json={"config": changed}).status_code == 200
    template = session.get(ExportTemplate, identifier)
    assert template.published_config == snapshot
    assert client.post(url + "/publish", json={"expected_version": 0}).status_code == 409
    assert client.post(url + "/publish", json={"expected_version": 1}).json()["version"] == 2
    assert client.delete(url).status_code == 409
    current["user"].role = UserRole.USER
    visible = client.get("/api/export-templates").json()
    assert visible[0]["draft_config"] is None
    assert visible[0]["published_config"]["title"] == "下一版"
    assert client.put(url, json={"enabled": False}).status_code == 403
    current["user"].role = UserRole.ADMIN
    assert client.put(url, json={"enabled": False}).status_code == 200
    current["user"].role = UserRole.USER
    assert client.get("/api/export-templates").json() == []
